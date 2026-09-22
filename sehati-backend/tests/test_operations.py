"""Import integrity tests. Every write uses the fixture's disposable PostgreSQL DB."""
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from test_api import environment  # Shared fixture creates/drops a random isolated DB.
from app import operations
from app.db import connect
from app.import_engine import DELIVERY_HEADERS
from app.workbook import HEADERS


def customer(code='TEST'):
    return operations.create_customer({'code': code, 'name': 'Import integrity test'})['id']


def legacy(tmp_path, name='legacy.xlsx'):
    """Two deliberately repeated D/Os, nullable groups and one standalone charge."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'January'
    sheet['C3'] = '01 Januari s.d. 31 Januari 2025'
    sheet['C4'] = 'INVOICE-TEST-001'
    for col, value in enumerate(HEADERS, 1):
        sheet.cell(6, col, value)
    rows = [
        [1, 'Test transporter', 'REPEATED-DO', 'TEST-1', date(2025, 1, 2), date(2025, 1, 3), date(2025, 1, 3), 'A', 'B', 'Customer', 45, 'Box', 10, 2.5, 100, 20, '=SUM(O7:P7)', 'first', None],
        [2, None, 'REPEATED-DO', 'TEST-2', date(2025, 1, 4), None, None, 'A', 'C', 'Customer', None, None, 20, None, None, None, None, 'second', None],
        [None] * 15 + [15, '=SUM(O9:P9)', 'invoice-level adjustment', None],
    ]
    for row_no, values in enumerate(rows, 7):
        for col, value in enumerate(values, 1):
            sheet.cell(row_no, col, value)
    path = tmp_path / name
    book.save(path)
    return path


def preview(path, cid):
    return operations.stage_import(path, path.name, cid)


def commit(path, cid):
    report = preview(path, cid)
    assert report['can_confirm'], report['issues']
    return operations.confirm_import(report['batch_id'], expected_revision=report['revision'])


def export(tmp_path, cid, name='export.xlsx', mutate=None):
    path = tmp_path / name
    payload = operations.export_workbook(cid)
    if mutate:
        book = openpyxl.load_workbook(BytesIO(payload))
        mutate(book)
        book.save(path)
        book.close()
    else:
        path.write_bytes(payload)
    return path


def set_field(book, field, value, row=2):
    book['Deliveries'].cell(row, DELIVERY_HEADERS.index(field) + 1, value)


def counts(cid):
    with connect() as conn:
        result = {}
        for table in ('deliveries', 'invoice_lines', 'billing_documents'):
            result[table] = conn.execute(f'SELECT count(*) AS n FROM {table} WHERE customer_id=%s', (cid,)).fetchone()['n']
        result['data_version'] = conn.execute('SELECT data_version FROM customers WHERE id=%s', (cid,)).fetchone()['data_version']
        return result


def test_real_legacy_preserves_all_source_rows_duplicates_and_invoice_charge(environment):
    path = Path('C:/Users/PC/Downloads/BRIDGESTONE 2025.xlsx')
    if not path.exists():
        pytest.skip('Private source workbook is not present')
    cid = customer()
    report = commit(path, cid)
    assert report['summary']['do_rows'] == 826
    assert report['summary']['distinct_do'] == 819
    assert report['total_stage_rows'] == 827
    with connect() as conn:
        assert conn.execute('SELECT count(*) AS n FROM deliveries').fetchone()['n'] == 826
        assert conn.execute('SELECT count(DISTINCT do_number) AS n FROM deliveries').fetchone()['n'] == 819
        assert conn.execute('SELECT count(*) AS n FROM billing_documents').fetchone()['n'] == 9
        assert conn.execute('SELECT count(*) AS n FROM invoice_lines WHERE delivery_id IS NULL').fetchone()['n'] == 1
        assert conn.execute('SELECT count(*) AS n FROM deliveries WHERE transporter IS NULL').fetchone()['n'] == 1
        assert conn.execute('SELECT count(*) AS n FROM import_rows').fetchone()['n'] == 827


def test_unchanged_export_keeps_versions_and_decimal_values(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)
    before = counts(cid)
    report = preview(export(tmp_path, cid), cid)
    assert report['summary']['unchanged'] == 3
    assert report['summary']['updated'] == 0
    committed = operations.confirm_import(report['batch_id'])
    assert committed['summary']['unchanged'] == 3
    assert counts(cid) == before
    with connect() as conn:
        assert {r['version'] for r in conn.execute('SELECT version FROM deliveries')} == {1}
        assert {r['version'] for r in conn.execute('SELECT version FROM invoice_lines')} == {1}


def test_edited_export_updates_only_target_preserving_ids(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)
    with connect() as conn:
        ids_before = {str(r['id']) for r in conn.execute('SELECT id FROM deliveries')}
    path = export(tmp_path, cid, mutate=lambda b: set_field(b, 'destination', 'Revised destination'))
    report = preview(path, cid)
    assert report['summary']['updated'] == 1
    assert report['summary']['new'] == 0
    operations.confirm_import(report['batch_id'])
    with connect() as conn:
        rows = conn.execute('SELECT id,version,destination FROM deliveries').fetchall()
        assert {str(r['id']) for r in rows} == ids_before
        assert len([r for r in rows if r['destination'] == 'Revised destination' and r['version'] == 2]) == 1
        assert len([r for r in rows if r['version'] == 1]) == 1


def test_missing_upload_rows_never_delete_deliveries_or_charges(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)
    before = counts(cid)

    def remove_rows(book):
        book['Deliveries'].delete_rows(2)
        book['Charges'].delete_rows(2, 10)

    report = commit(export(tmp_path, cid, mutate=remove_rows), cid)
    assert report['summary']['do_rows'] == 1
    assert counts(cid) == before


def test_old_offline_version_blocked_before_confirmation(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)
    stale = export(tmp_path, cid, 'stale.xlsx', lambda b: set_field(b, 'destination', 'Stale edit'))
    fresh = export(tmp_path, cid, 'fresh.xlsx', lambda b: set_field(b, 'destination', 'Newer edit'))
    commit(fresh, cid)
    before = counts(cid)
    report = preview(stale, cid)
    assert not report['can_confirm']
    assert any(i['code'] == 'STALE_OR_FOREIGN_RECORD' for i in report['issues'])
    with pytest.raises(operations.OperationsError) as raised:
        operations.confirm_import(report['batch_id'])
    assert raised.value.status_code == 409
    assert counts(cid) == before


def test_database_change_after_preview_rolls_back_and_retains_failure_audit(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)
    first = preview(export(tmp_path, cid, 'first.xlsx', lambda b: set_field(b, 'destination', 'First editor')), cid)
    second = preview(export(tmp_path, cid, 'second.xlsx', lambda b: set_field(b, 'destination', 'Second editor')), cid)
    operations.confirm_import(first['batch_id'])
    before = counts(cid)
    with pytest.raises(operations.OperationsError):
        operations.confirm_import(second['batch_id'])
    assert counts(cid) == before
    assert operations.get_batch(second['batch_id'])['status'] == 'failed'
    with connect() as conn:
        assert conn.execute('SELECT count(*) AS n FROM deliveries WHERE destination=%s', ('Second editor',)).fetchone()['n'] == 0
        assert conn.execute("SELECT count(*) AS n FROM audit_events WHERE action='import_confirm_failed' AND object_id=%s", (second['batch_id'],)).fetchone()['n'] == 1


def test_foreign_customer_ids_blocked_even_if_metadata_relabelled(environment, tmp_path):
    first, second = customer('FIRST'), customer('SECOND')
    commit(legacy(tmp_path), first)

    def relabel(book):
        book['Metadata']['B2'] = second

    report = preview(export(tmp_path, first, mutate=relabel), second)
    assert not report['can_confirm']
    assert any(i['code'] == 'STALE_OR_FOREIGN_RECORD' for i in report['issues'])
    assert counts(second)['deliveries'] == 0


def test_legacy_replay_never_reverts_committed_stable_id_edits(environment, tmp_path):
    cid = customer()
    source = legacy(tmp_path)
    initial = commit(source, cid)
    commit(export(tmp_path, cid, mutate=lambda b: set_field(b, 'destination', 'Corrected')), cid)
    before = counts(cid)
    report = preview(source, cid)
    assert report['reused']
    assert report['batch_id'] == initial['batch_id']
    operations.confirm_import(report['batch_id'])
    assert counts(cid) == before
    with connect() as conn:
        assert conn.execute('SELECT count(*) AS n FROM deliveries WHERE destination=%s', ('Corrected',)).fetchone()['n'] == 1


def test_resaved_legacy_with_unchanged_content_is_noop_after_later_correction(environment, tmp_path):
    cid = customer()
    source = legacy(tmp_path)
    commit(source, cid)
    commit(export(tmp_path, cid, mutate=lambda b: set_field(b, 'destination', 'Corrected')), cid)
    before = counts(cid)
    book = openpyxl.load_workbook(source)
    book.properties.creator = 'Another spreadsheet application'
    resaved = tmp_path / 'resaved.xlsx'
    book.save(resaved)
    book.close()
    report = preview(resaved, cid)
    assert report['reused'] and report['status'] == 'committed'
    operations.confirm_import(report['batch_id'])
    assert counts(cid) == before


def test_changed_legacy_requires_stable_id_export(environment, tmp_path):
    cid = customer()
    source = legacy(tmp_path)
    commit(source, cid)
    book = openpyxl.load_workbook(source)
    book.active['I7'] = 'Unmatched manual correction'
    changed = tmp_path / 'changed.xlsx'
    book.save(changed)
    book.close()
    report = preview(changed, cid)
    assert not report['can_confirm']
    assert any(i['code'] == 'LEGACY_UPDATE_REQUIRES_IDS' for i in report['issues'])


def test_confirm_is_idempotent_and_has_one_commit_audit(environment, tmp_path):
    cid = customer()
    report = commit(legacy(tmp_path), cid)
    before = counts(cid)
    again = operations.confirm_import(report['batch_id'])
    assert again['reused']
    assert counts(cid) == before
    with connect() as conn:
        assert conn.execute("SELECT count(*) AS n FROM audit_events WHERE action='import_committed' AND object_id=%s", (report['batch_id'],)).fetchone()['n'] == 1


def test_transaction_failure_commits_neither_records_nor_import_rows(environment, tmp_path, monkeypatch):
    cid = customer()
    report = preview(legacy(tmp_path), cid)
    original = operations._write
    calls = []

    def fail_second_write(*args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            raise RuntimeError('Injected interruption')
        return original(*args, **kwargs)

    monkeypatch.setattr(operations, '_write', fail_second_write)
    with pytest.raises(operations.OperationsError):
        operations.confirm_import(report['batch_id'])
    assert len(calls) == 2
    assert counts(cid) == {'deliveries': 0, 'invoice_lines': 0, 'billing_documents': 0, 'data_version': 0}
    failed = operations.get_batch(report['batch_id'])
    assert failed['status'] == 'failed'
    assert failed['summary']['issue_count'] == len(failed['issues'])
    with connect() as conn:
        assert conn.execute('SELECT count(*) AS n FROM import_rows').fetchone()['n'] == 0
        assert conn.execute('SELECT count(*) AS n FROM legacy_document_snapshots').fetchone()['n'] == 0
        assert conn.execute("SELECT count(*) AS n FROM audit_events WHERE action='import_confirm_failed'").fetchone()['n'] == 1


def test_two_previews_of_same_new_records_do_not_double_insert(environment, tmp_path):
    cid = customer()

    def add_delivery(book):
        values = {'invoice_number': 'NEW-001', 'period_start': '2025-01-01', 'period_end': '2025-01-31', 'do_number': 'NEW-DO', 'vehicle_plate': 'TEST-1', 'load_date': '2025-01-02'}
        book['Deliveries'].append([values.get(h) for h in DELIVERY_HEADERS])

    source = export(tmp_path, cid, mutate=add_delivery)
    first, second = preview(source, cid), preview(source, cid)
    assert first['batch_id'] != second['batch_id']
    operations.confirm_import(first['batch_id'])
    before = counts(cid)
    operations.confirm_import(second['batch_id'])
    assert counts(cid) == before
    assert before['deliveries'] == 1
    replayed_rows = operations.get_batch_rows(second['batch_id'], limit=1)['items']
    original_rows = operations.get_batch_rows(first['batch_id'], limit=1)['items']
    assert replayed_rows[0]['record_id'] == original_rows[0]['record_id']
    assert replayed_rows[0]['action'] == 'unchanged'
    assert replayed_rows[0]['changes'] == []
    with connect() as conn:
        assert conn.execute('SELECT 1 FROM deliveries WHERE id=%s AND customer_id=%s', (replayed_rows[0]['record_id'], cid)).fetchone()


def test_semantically_identical_new_template_previews_do_not_double_insert(environment, tmp_path):
    cid = customer()

    def add_delivery(book):
        values = {'invoice_number': 'NEW-001', 'period_start': '2025-01-01', 'period_end': '2025-01-31', 'do_number': 'NEW-DO', 'vehicle_plate': 'TEST-1', 'load_date': '2025-01-02'}
        book['Deliveries'].append([values.get(h) for h in DELIVERY_HEADERS])

    source = export(tmp_path, cid, mutate=add_delivery)
    book = openpyxl.load_workbook(source)
    book.properties.creator = 'Resaved in another application'
    resaved = tmp_path / 'resaved-new-record.xlsx'
    book.save(resaved)
    book.close()
    assert source.read_bytes() != resaved.read_bytes()
    first, second = preview(source, cid), preview(resaved, cid)
    operations.confirm_import(first['batch_id'])
    before = counts(cid)
    report = operations.confirm_import(second['batch_id'])
    assert report['reused']
    assert counts(cid) == before
    assert before['deliveries'] == 1


def test_unknown_nonempty_main_row_blocks_batch(environment, tmp_path):
    cid = customer()
    source = legacy(tmp_path)
    book = openpyxl.load_workbook(source)
    book.active['R12'] = 'Unclassified note must not disappear'
    book.save(source)
    book.close()
    report = preview(source, cid)
    assert not report['can_confirm']
    assert any(i['severity'] == 'blocking' and i.get('cells') for i in report['issues'])


def test_far_right_extra_column_in_stable_export_blocks_batch(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)

    def add_unmapped_value(book):
        book['Deliveries'].cell(2, len(DELIVERY_HEADERS) + 3, 'Unmapped data')

    report = preview(export(tmp_path, cid, mutate=add_unmapped_value), cid)
    assert not report['can_confirm']
    assert any(i['severity'] == 'blocking' for i in report['issues'])


def test_hidden_stable_sheet_blocks_batch(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)
    before = counts(cid)

    def hide_deliveries(book):
        book['Deliveries'].sheet_state = 'hidden'

    report = preview(export(tmp_path, cid, mutate=hide_deliveries), cid)
    assert not report['can_confirm']
    assert any(i['code'] == 'HIDDEN_SHEET' and i['sheet'] == 'Deliveries' and i['severity'] == 'blocking' for i in report['issues'])
    with pytest.raises(operations.OperationsError):
        operations.confirm_import(report['batch_id'])
    assert counts(cid) == before


def test_charge_version_without_charge_identity_is_invalid(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)

    def orphan_charge_version(book):
        sheet = book['Deliveries']
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, DELIVERY_HEADERS.index('charge_id') + 1).value is None:
                set_field(book, 'charge_version', 1, row)
                return
        raise AssertionError('Fixture requires an uncharged delivery')

    report = preview(export(tmp_path, cid, mutate=orphan_charge_version), cid)
    assert not report['can_confirm']


def test_export_escapes_formula_like_identifiers_as_text(environment, tmp_path):
    cid = customer()
    commit(legacy(tmp_path), cid)
    with connect() as conn:
        conn.execute('UPDATE deliveries SET do_number=%s,notes=%s WHERE customer_id=%s', ('=HYPERLINK("https://example.invalid")', '+CMD', cid))
    book = openpyxl.load_workbook(BytesIO(operations.export_workbook(cid)), data_only=False)
    cell = book['Deliveries'].cell(2, DELIVERY_HEADERS.index('do_number') + 1)
    assert cell.data_type == 's' and cell.value.startswith('=')
    book.close()
