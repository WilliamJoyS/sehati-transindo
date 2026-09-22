"""Bounded, read-only review of the supplied Bridgestone XLSX layout.

This module never executes formulas, fills down blanks, infers trips, writes the
source file, or enables a database import. Business mappings remain unapproved.
Reports contain coordinates and aggregate counts, never D/O IDs or money values.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path, PurePosixPath
import math
import re
from defusedxml import ElementTree as ET
from defusedxml.common import DefusedXmlException
from zipfile import BadZipFile, ZipFile

import openpyxl
from openpyxl.utils.cell import column_index_from_string, get_column_letter

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 2_000
MAX_XML_BYTES = 30 * 1024 * 1024
MAX_POPULATED_CELLS = 200_000
MAX_PHYSICAL_CELLS = 300_000
MAX_SHEETS = 24
MAX_BUSINESS_ROWS = 10_000
MAX_ROW_INDEX = 25_000
MAX_COLUMN_INDEX = 64
MAX_ITERATION_CELLS = 1_000_000
MAPPING_VERSION = "bridgestone-review-v1"

HEADERS = (
    "No.", "TRANSPORTER", "No. D/O", "No.  TRUCK", "Tanggal Muat",
    "Perkiraan Tgl Bongkar", "Actual Tgl Bongkar", "Dari", "Tujuan",
    "Agen/ Distributor", "Kapasitas Truck (M3) - (45/60)", "Tipe",
    "Qty tire", "M3", "Tarif Normal", "Tarif Tambahan (abnormal)",
    "Total Tagihan", "Ket.", "Rit",
)
MONTHS = dict(zip(
    ("januari", "februari", "maret", "april", "mei", "juni", "juli",
     "agustus", "september", "oktober", "november", "desember"),
    range(1, 13),
))
PERIOD_RE = re.compile(
    r"\(?\s*(\d{1,2})\s+([A-Za-z]+)\s+s\s*\.\s*d\s*\.\s*"
    r"(\d{1,2})\s+([A-Za-z]+)\s*\)?\s*(\d{4})", re.I,
)
SUMMARY_LABELS = {
    "TOTAL TAGIHAN REP BEKASI", "TOTAL TAGIHAN REP KARAWANG",
    "TOTAL TAGIHAN BEKASI & KARAWANG", "DPP LAIN", "PPN", "GRAND TOTAL",
}


class WorkbookValidationError(ValueError):
    """The upload cannot safely be opened or exceeds supported limits."""


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _normalize(value):
    if _blank(value):
        return ""
    return " ".join(str(value).split()).casefold()


def _xml_root(raw: bytes):
    # Reject DTD/entity declarations rather than accepting untrusted expansions.
    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise WorkbookValidationError("XML document types and entities are unsupported.")
    try:
        return ET.fromstring(raw)
    except (ET.ParseError, DefusedXmlException) as exc:
        raise WorkbookValidationError("The workbook contains malformed XML.") from exc


def _preflight(path: Path):
    """Check the archive before openpyxl reads it; do not extract any entries."""
    if not path.is_file() or path.stat().st_size == 0:
        raise WorkbookValidationError("The upload is empty or unavailable.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise WorkbookValidationError("The XLSX upload exceeds the 10 MB limit.")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise WorkbookValidationError("The workbook contains too many archive entries.")
            if sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                raise WorkbookValidationError("Expanded workbook data exceeds 100 MB.")
            names = set()
            populated = physical = iteration_cells = 0
            worksheet_count = 0
            for item in entries:
                name = item.filename
                canonical = name.casefold().replace("\\", "/")
                parts = PurePosixPath(canonical).parts
                if (canonical.startswith("/") or ".." in parts or ":" in canonical
                        or "\\" in name or canonical in names):
                    raise WorkbookValidationError("Unsafe or duplicate archive entry names.")
                names.add(canonical)
                if item.flag_bits & 1:
                    raise WorkbookValidationError("Encrypted workbooks are unsupported.")
                if (canonical.endswith(".bin") or "vbaproject" in canonical
                        or canonical.startswith(("xl/externallinks/", "xl/activex/", "xl/embeddings/"))):
                    raise WorkbookValidationError("Macros, embedded objects, and external workbook links are unsupported.")
                if not canonical.endswith((".xml", ".rels")):
                    continue
                if item.file_size > MAX_XML_BYTES:
                    raise WorkbookValidationError("An XML workbook part exceeds the supported size.")
                root = _xml_root(archive.read(item))
                if canonical.endswith(".rels"):
                    if any(node.attrib.get("TargetMode", "").casefold() == "external" for node in root):
                        raise WorkbookValidationError("External relationships are unsupported in uploads.")
                if canonical == "[content_types].xml":
                    if any("macroenabled" in node.attrib.get("ContentType", "").casefold() for node in root):
                        raise WorkbookValidationError("Macro-enabled workbook content is unsupported.")
                if canonical == "xl/workbook.xml":
                    sheet_nodes = [node for node in root.iter() if node.tag.rsplit("}", 1)[-1] == "sheet"]
                    if not 1 <= len(sheet_nodes) <= MAX_SHEETS:
                        raise WorkbookValidationError("The workbook must have between 1 and 24 worksheets.")
                if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", canonical):
                    continue
                worksheet_count += 1
                if worksheet_count > MAX_SHEETS:
                    raise WorkbookValidationError("The workbook exceeds 24 worksheets.")
                max_row = max_col = 0
                for node in root.iter():
                    if node.tag.rsplit("}", 1)[-1] != "c":
                        continue
                    physical += 1
                    coordinate = re.fullmatch(r"([A-Za-z]+)([1-9]\d*)", node.attrib.get("r", ""))
                    if not coordinate:
                        raise WorkbookValidationError("A cell has an invalid coordinate.")
                    column = column_index_from_string(coordinate[1])
                    row = int(coordinate[2])
                    if row > MAX_ROW_INDEX or column > MAX_COLUMN_INDEX:
                        raise WorkbookValidationError("The worksheet dimensions exceed supported limits.")
                    max_row, max_col = max(row, max_row), max(column, max_col)
                    if any(child.tag.rsplit("}", 1)[-1] in {"v", "f", "is"} for child in node):
                        populated += 1
                    if populated > MAX_POPULATED_CELLS or physical > MAX_PHYSICAL_CELLS:
                        raise WorkbookValidationError("The workbook exceeds supported cell limits.")
                iteration_cells += max_row * max(19, max_col)
                if iteration_cells > MAX_ITERATION_CELLS:
                    raise WorkbookValidationError("Sparse worksheet dimensions exceed supported processing limits.")
            if not {"[content_types].xml", "xl/workbook.xml", "_rels/.rels"}.issubset(names):
                raise WorkbookValidationError("The file is not a supported XLSX workbook.")
            if not worksheet_count:
                raise WorkbookValidationError("No supported worksheets were found.")
    except WorkbookValidationError:
        raise
    except (BadZipFile, OSError, ValueError, RuntimeError, KeyError, NotImplementedError) as exc:
        raise WorkbookValidationError("The file is not a readable, unencrypted XLSX workbook.") from exc


def _period(value):
    match = PERIOD_RE.fullmatch(str(value or "").strip().lstrip(": "))
    if not match:
        return None
    first, first_month, last, last_month, year = match.groups()
    try:
        start = date(int(year), MONTHS[first_month.casefold()], int(first))
        end = date(int(year), MONTHS[last_month.casefold()], int(last))
    except (KeyError, ValueError):
        return None
    if end < start:
        return None
    label = f"{start.day:02d} {first_month.capitalize()} – {end.day:02d} {last_month.capitalize()} {year}"
    return label, start, end


def _cell_fingerprint(value):
    # Only used to compare repeated IDs; raw financial values never leave here.
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    return sha256(repr(value).encode("utf-8")).hexdigest()


def analyze_workbook(path: Path, filename: str | None = None) -> dict:
    path = Path(path)
    filename = (filename or path.name).replace("\\", "/").rsplit("/", 1)[-1][:255]
    if Path(filename).suffix.casefold() != ".xlsx":
        raise WorkbookValidationError("Only .xlsx files are supported.")
    _preflight(path)
    issues = []

    def issue(severity, code, message, sheet=None, cells=None):
        item = {"severity": severity, "code": code, "message": message}
        if sheet is not None:
            item["sheet"] = sheet
        if cells:
            item["cells"] = list(cells[:20])
            if len(cells) > 20:
                item["additional_cell_count"] = len(cells) - 20
        issues.append(item)

    issue("blocking", "BUSINESS_RULES_PENDING",
          "Preview only: stable identifiers, grouped rows, charges, and correction rules require approval before any import can be committed.")
    book = cached = None
    try:
        # data_only retrieves stored Excel caches; neither open executes formulas.
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        if len(book.worksheets) != len(book.sheetnames):
            raise WorkbookValidationError("Non-worksheet sheet types are unsupported.")
        periods = []
        identifiers = defaultdict(list)
        total_rows = 0
        for sheet in book:
            if sheet.sheet_state != "visible":
                issue("blocking", "HIDDEN_SHEET", "A hidden worksheet requires explicit review.", sheet.title)
            # Ignore supplied dimension metadata: it may be stale or malicious.
            sheet.reset_dimensions()
            cached_sheet = cached[sheet.title]
            cached_sheet.reset_dimensions()
            rows = list(sheet.iter_rows(max_col=MAX_COLUMN_INDEX))
            cached_rows = list(cached_sheet.iter_rows(max_col=MAX_COLUMN_INDEX))
            def cell(row, col, cache=False):
                target = cached_rows if cache else rows
                return target[row - 1][col - 1] if row <= len(target) else None
            def value(row, col, cache=False):
                item = cell(row, col, cache)
                return item.value if item is not None else None

            observed = [_normalize(value(6, col)) for col in range(1, 20)]
            expected = [_normalize(header) for header in HEADERS]
            duplicate_headers = [key for key, count in Counter(observed).items() if key and count > 1]
            if duplicate_headers:
                issue("blocking", "DUPLICATE_HEADERS", "Duplicate column labels make mapping ambiguous.", sheet.title, ["A6:S6"])
            if observed != expected:
                mismatches = [f"{get_column_letter(index + 1)}6" for index, pair in enumerate(zip(observed, expected)) if pair[0] != pair[1]]
                issue("blocking", "HEADER_MISMATCH", "Row 6 does not match the reviewed 19-column Bridgestone layout; rows were not interpreted.", sheet.title, mismatches)
                periods.append({"sheet": sheet.title, "period": "Unrecognized layout", "do_rows": 0})
                continue
            period = _period(value(3, 3))
            entry = {"sheet": sheet.title, "period": period[0] if period else "Unrecognized billing period", "do_rows": 0}
            if period:
                entry.update(period_start=period[1].isoformat(), period_end=period[2].isoformat())
            else:
                issue("blocking", "INVALID_PERIOD", "C3 must contain a recognized billing date range.", sheet.title, ["C3"])
            if _blank(value(4, 3)):
                issue("warning", "MISSING_INVOICE_REFERENCE", "The invoice reference is missing; its rules still need agreement.", sheet.title, ["C4"])
            missing = defaultdict(list)
            grouped = defaultdict(list)
            invalid = defaultdict(list)
            formula_cells = []
            summary_rows = []
            auxiliary_rows = []
            blank_rit = 0
            date_outside = []
            for row_index, row in enumerate(rows, start=1):
                if row_index < 7:
                    continue
                main = [item.value for item in row[:19]]
                if any(not _blank(item.value) for item in row[19:]):
                    auxiliary_rows.append(f"T{row_index}:BL{row_index}")
                if all(_blank(item) for item in main):
                    continue
                if isinstance(main[0], str) and " ".join(main[0].upper().split()) in SUMMARY_LABELS:
                    summary_rows.append(f"A{row_index}:S{row_index}")
                    continue
                if [_normalize(item) for item in main] == expected:
                    issue("blocking", "REPEATED_HEADER_ROW", "A repeated header inside the data region requires review.", sheet.title, [f"A{row_index}:S{row_index}"])
                    continue
                if all(_blank(main[index]) for index in range(14)) and any(not _blank(main[index]) for index in (14, 15, 16)):
                    issue("blocking", "UNLINKED_CHARGE", "A charge row has no D/O, truck, or loading date. Its invoice/trip/delivery linkage must be agreed; it was not discarded.", sheet.title, [f"O{row_index}:R{row_index}"])
                    continue
                if all(_blank(main[index]) for index in (2, 3, 4)):
                    issue("blocking", "UNCLASSIFIED_ROW", "A populated row does not match a delivery, approved summary, or standalone charge. Manual classification is required.", sheet.title, [f"A{row_index}:S{row_index}"])
                    continue
                if _blank(main[2]):
                    missing["D/O"].append(f"C{row_index}")
                    continue
                entry["do_rows"] += 1
                total_rows += 1
                if total_rows > MAX_BUSINESS_ROWS:
                    raise WorkbookValidationError("The workbook exceeds 10,000 business rows.")
                identifier_cell = row[2]
                identifier = main[2]
                if identifier_cell.data_type == "f":
                    invalid["formula identifier"].append(f"C{row_index}")
                else:
                    if isinstance(identifier, (int, float)) and not isinstance(identifier, bool):
                        invalid["numeric identifier"].append(f"C{row_index}")
                        if isinstance(identifier, float) and identifier.is_integer():
                            identifier = int(identifier)
                    identifiers[str(identifier).strip().upper()].append((
                        sheet.title, row_index,
                        tuple(_cell_fingerprint(value(row_index, col, True)) for col in range(2, 19)),
                    ))
                for col, field in ((2, "transporter"), (4, "truck"), (5, "loading date"), (14, "volume")):
                    if _blank(main[col - 1]):
                        missing[field].append(f"{get_column_letter(col)}{row_index}")
                for col in (11, 12, 15, 16, 17):
                    if _blank(main[col - 1]):
                        grouped[get_column_letter(col)].append(f"{get_column_letter(col)}{row_index}")
                blank_rit += int(_blank(main[18]))
                for col in (5, 6, 7):
                    current = main[col - 1]
                    if not _blank(current) and not isinstance(current, (datetime, date)):
                        invalid["date type"].append(f"{get_column_letter(col)}{row_index}")
                loaded = main[4]
                if isinstance(loaded, datetime):
                    loaded = loaded.date()
                if period and isinstance(loaded, date) and not period[1] <= loaded <= period[2]:
                    date_outside.append(f"E{row_index}")
                for col in (11, 13, 14, 15, 16, 17, 19):
                    current = main[col - 1]
                    current_cell = row[col - 1]
                    if current_cell.data_type == "f":
                        formula_cells.append(f"{get_column_letter(col)}{row_index}")
                        continue
                    if _blank(current):
                        continue
                    if not isinstance(current, (int, float)) or isinstance(current, bool) or not math.isfinite(current):
                        invalid["numeric type"].append(f"{get_column_letter(col)}{row_index}")
                    elif col in (13, 19) and current != int(current):
                        invalid["whole-number quantity"].append(f"{get_column_letter(col)}{row_index}")
            for field, coordinates in missing.items():
                severity = "blocking" if field in {"D/O", "truck", "loading date"} else "warning"
                issue(severity, "MISSING_" + field.upper().replace("/", "").replace(" ", "_"),
                      f"{len(coordinates)} row(s) have a blank {field}; no value was inferred.", sheet.title, coordinates)
            for kind, coordinates in invalid.items():
                issue("warning" if kind == "numeric identifier" else "blocking",
                      kind.upper().replace(" ", "_"),
                      f"{len(coordinates)} cell(s) require review for {kind}. IDs are treated as text; lost leading zeros cannot be reconstructed.", sheet.title, coordinates)
            if grouped:
                counts = ", ".join(f"{column}: {len(coordinates)}" for column, coordinates in grouped.items())
                issue("warning", "GROUPED_BLANKS_UNRESOLVED",
                      f"Blank vehicle/charge fields ({counts}) may belong to grouped rows. No fill-down, zero substitution, or charge allocation was applied.", sheet.title)
            if blank_rit:
                issue("info", "BLANK_RIT", f"Rit is blank on {blank_rit} D/O row(s); its meaning and requirement are unconfirmed.", sheet.title)
            if formula_cells:
                issue("info", "FORMULAS_NOT_EXECUTED", f"{len(formula_cells)} data cells contain formulas. Stored caches are used only for comparison; formulas were not executed or financially validated.", sheet.title, formula_cells)
            if summary_rows:
                issue("info", "SUMMARY_ROWS_EXCLUDED", f"{len(summary_rows)} subtotal/tax/total rows were classified separately from delivery rows. Reconciliation and tax rules require approval.", sheet.title, summary_rows)
            if auxiliary_rows:
                issue("info", "AUXILIARY_DATA_EXCLUDED", f"{len(auxiliary_rows)} row(s) contain reference data outside A:S. These cells were not imported as deliveries or tariff rules.", sheet.title, auxiliary_rows)
            if date_outside:
                issue("warning", "LOADING_DATE_OUTSIDE_PERIOD", f"{len(date_outside)} loading date(s) fall outside the sheet billing period; billing rules may explain this.", sheet.title, date_outside)
            periods.append(entry)

        for occurrences in identifiers.values():
            if len(occurrences) < 2:
                continue
            different = [get_column_letter(index + 2) for index in range(17)
                         if len({item[2][index] for item in occurrences}) > 1]
            references = [f"{name}!C{row}" for name, row, _ in occurrences]
            detail = ("Different stored fields occur in columns " + ", ".join(different) + ".") if different else "The stored business fields match."
            issue("blocking", "REPEATED_DO", f"One D/O value appears in {len(occurrences)} rows. {detail} Review split deliveries, corrections, or duplicate-entry rules before selecting an import key.", cells=references)
        ordered_periods = sorted((item for item in periods if "period_start" in item), key=lambda item: item["period_start"])
        if any(current["period_start"] <= previous["period_end"] for previous, current in zip(ordered_periods, ordered_periods[1:])):
            issue("info", "OVERLAPPING_PERIODS", "Billing periods overlap. Month names, sheet names, and date ranges cannot serve as unique delivery keys.")
        if not total_rows:
            issue("blocking", "NO_DELIVERY_ROWS", "No D/O rows were recognized in the reviewed layout.")
        return {
            "source": filename,
            "source_metadata": {"size_bytes": path.stat().st_size, "sha256": sha256(path.read_bytes()).hexdigest(), "mapping_version": MAPPING_VERSION},
            "summary": {"sheets": len(book.worksheets), "do_rows": total_rows, "distinct_do": len(identifiers), "issue_count": len(issues)},
            "periods": periods,
            "issues": issues,
            "commit_enabled": False,
            "status": "preview_only",
        }
    except WorkbookValidationError:
        raise
    except Exception as exc:
        # The API must not leak workbook values through parser exception text.
        raise WorkbookValidationError("The workbook could not be read using the supported layout.") from exc
    finally:
        if book is not None:
            book.close()
        if cached is not None:
            cached.close()
