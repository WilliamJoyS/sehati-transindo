"""Read approved layouts without executing Excel formulas or guessing grouping."""
from __future__ import annotations
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
from uuid import UUID, uuid4
import json
import re
import openpyxl
from .workbook import (_preflight, _normalize, _blank, _period, HEADERS, SUMMARY_LABELS,
                       WorkbookValidationError, MAX_BUSINESS_ROWS)

VERSION = 'sehati-records-v1'
LEGACY_VERSION = 'bridgestone-operational-v1'
DELIVERY_FIELDS = ('do_number','vehicle_plate','transporter','load_date','estimated_unload_date',
                   'actual_unload_date','origin','destination','distributor','capacity_m3',
                   'vehicle_type','quantity','volume_m3','notes')
CHARGE_FIELDS = ('normal_charge','additional_charge','reported_total','note')
DOCUMENT_FIELDS = ('invoice_number','period_start','period_end')
DELIVERY_HEADERS = ('record_id','record_version') + DOCUMENT_FIELDS + DELIVERY_FIELDS + ('charge_id','charge_version') + CHARGE_FIELDS
CHARGE_HEADERS = ('charge_id','charge_version') + DOCUMENT_FIELDS + CHARGE_FIELDS

def clean(value):
    if isinstance(value, datetime): return value.isoformat()
    if isinstance(value, date): return value.isoformat()
    if isinstance(value, Decimal): return format(value.normalize(),'f')
    if isinstance(value, UUID): return str(value)
    if isinstance(value, dict): return {k:clean(v) for k,v in value.items()}
    if isinstance(value, (tuple,list)): return [clean(v) for v in value]
    return value

def digest(value):
    return sha256(json.dumps(clean(value),sort_keys=True,separators=(',',':'),ensure_ascii=False).encode()).hexdigest()

def issue(code,message,sheet=None,cells=None,severity='blocking'):
    return {'code':code,'message':message,'sheet':sheet,'cells':cells or [],'severity':severity}

def text_value(value, required=False):
    if _blank(value):
        if required: raise ValueError('Required value is blank.')
        return None
    if isinstance(value,bool): raise ValueError('Boolean values are not valid identifiers or text.')
    if isinstance(value,float) and value.is_integer(): value=int(value)
    value=str(value).strip()
    if len(value)>2000: raise ValueError('Text exceeds 2,000 characters.')
    return value

def number(value,integer=False,money=False):
    if _blank(value): return None
    if isinstance(value,bool): raise ValueError('Boolean is not a number.')
    try: parsed=Decimal(str(value))
    except InvalidOperation: raise ValueError('Invalid number.')
    if not parsed.is_finite() or abs(parsed)>=Decimal('1000000000000'): raise ValueError('Number is outside supported limits.')
    if integer:
        if parsed!=parsed.to_integral_value() or parsed<0: raise ValueError('Quantity must be a nonnegative whole number.')
        return int(parsed)
    places=Decimal('0.01') if money else Decimal('0.000001')
    if parsed!=parsed.quantize(places): raise ValueError('Too many decimal places.')
    if not money and parsed<0: raise ValueError('Capacity and volume cannot be negative.')
    return format(parsed.normalize(),'f')

def date_value(value,required=False):
    if _blank(value):
        if required: raise ValueError('Loading date is required.')
        return None
    if isinstance(value,datetime): return value.date().isoformat()
    if isinstance(value,date): return value.isoformat()
    try: return date.fromisoformat(str(value)).isoformat()
    except ValueError: raise ValueError('Dates must be Excel dates or YYYY-MM-DD.')

def record_id(value,version):
    if _blank(value):
        if not _blank(version): raise ValueError('A version without its record ID is invalid.')
        return str(uuid4()),None
    try: parsed=str(UUID(str(value)))
    except ValueError: raise ValueError('Record ID must be a UUID from the export.')
    parsed_version=number(version,integer=True)
    if not parsed_version: raise ValueError('Existing record ID requires its positive version.')
    return parsed,parsed_version

def document(values):
    result={'invoice_number':text_value(values['invoice_number'],True),
            'period_start':date_value(values['period_start'],True),'period_end':date_value(values['period_end'],True)}
    if result['period_start']>result['period_end']: raise ValueError('Billing period ends before it starts.')
    return result

def delivery(values):
    result={}
    for field in DELIVERY_FIELDS:
        value=values.get(field)
        if field.endswith('_date'): result[field]=date_value(value,field=='load_date')
        elif field in {'capacity_m3','volume_m3'}: result[field]=number(value)
        elif field=='quantity': result[field]=number(value,integer=True)
        else: result[field]=text_value(value,field in {'do_number','vehicle_plate'})
    if result['actual_unload_date'] and result['actual_unload_date']<result['load_date']:
        raise ValueError('Actual unload date precedes loading date.')
    return result

def charge(values):
    return {field:text_value(values.get(field)) if field=='note' else number(values.get(field),money=True) for field in CHARGE_FIELDS}

def read_workbook(path, customer_id):
    path=Path(path)
    if path.suffix.lower()!='.xlsx': raise WorkbookValidationError('Only .xlsx files are supported.')
    _preflight(path)
    book=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
    cached=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        result={'mapping_version':VERSION if 'Metadata' in book.sheetnames else LEGACY_VERSION,'rows':[], 'issues':[], 'periods':[], 'documents':{}}
        for sheet in book:
            if sheet.sheet_state!='visible':
                result['issues'].append(issue('HIDDEN_SHEET','Unhide all sheets and review them before importing.',sheet.title))
        if result['mapping_version']==VERSION:
            _read_export(book,customer_id,result)
        else:
            _read_legacy(book,cached,result)
        if len(result['rows'])>MAX_BUSINESS_ROWS: raise WorkbookValidationError('Maximum 10,000 business rows per upload.')
        if not result['rows']: result['issues'].append(issue('NO_ROWS','No business records were recognized.'))
        for row in result['rows']:
            name=row['document']['invoice_number']
            doc=result['documents'].setdefault(name,{**row['document'],'rows':[]})
            if any(doc[k]!=row['document'][k] for k in DOCUMENT_FIELDS):
                result['issues'].append(issue('INCONSISTENT_DOCUMENT','The same invoice reference has conflicting billing periods.',row['sheet'],[str(row['row_number'])]))
            doc['rows'].append(row)
        for doc in result['documents'].values():
            doc['content_digest']=digest(sorted(digest({'kind':r['kind'],'data':r['data'],'charge':(r.get('charge') or {}).get('data')}) for r in doc.pop('rows')))
        # Generated IDs are deliberately excluded for new rows. This identifies a
        # re-saved or reordered copy without mistaking it for a fresh insertion.
        result['content_digest']=digest({'version':result['mapping_version'], 'rows':sorted(
            digest({'kind':r['kind'],'document':r['document'],'data':r['data'],
                    'id':r['id'] if r['expected_version'] is not None else None,
                    'version':r['expected_version'],
                    'charge':{'data':r['charge']['data'],
                              'id':r['charge']['id'] if r['charge']['expected_version'] is not None else None,
                              'version':r['charge']['expected_version']} if r.get('charge') else None})
            for r in result['rows'])})
        return result
    except WorkbookValidationError: raise
    except Exception as exc:
        raise WorkbookValidationError('Workbook could not be read using a supported operational template.') from exc
    finally: book.close(); cached.close()

def _read_export(book,customer_id,result):
    if set(book.sheetnames)!={'Metadata','Deliveries','Charges'}:
        result['issues'].append(issue('SHEET_MISMATCH','Stable-ID workbook requires Metadata, Deliveries, and Charges sheets only.')); return
    meta={r[0]:r[1] for r in book['Metadata'].iter_rows(max_col=2,values_only=True) if r[0]}
    if meta.get('template_version')!=VERSION or str(meta.get('customer_id'))!=str(customer_id):
        result['issues'].append(issue('CUSTOMER_TEMPLATE_MISMATCH','Export belongs to a different customer or template version.')); return
    seen=set()
    for name,headers,kind in [('Deliveries',DELIVERY_HEADERS,'delivery'),('Charges',CHARGE_HEADERS,'charge')]:
        sheet=book[name]; sheet.reset_dimensions()
        rows=sheet.iter_rows(max_col=64)
        first=next(rows,[])
        if tuple(c.value for c in first[:len(headers)])!=headers or any(not _blank(c.value) for c in first[len(headers):]):
            result['issues'].append(issue('HEADER_MISMATCH','Export column names must remain unchanged.',name)); continue
        count=0
        for index,cells in enumerate(rows,2):
            values=[c.value for c in cells]
            if all(_blank(v) for v in values): continue
            count+=1
            try:
                if any(not _blank(v) for v in values[len(headers):]): raise ValueError('Unexpected extra column contains data.')
                if any(c.data_type=='f' for c in cells): raise ValueError('Export imports accept values only; replace formulas with reviewed values.')
                val=dict(zip(headers,values)); doc=document(val)
                ident,version=record_id(val['record_id' if kind=='delivery' else 'charge_id'],val['record_version' if kind=='delivery' else 'charge_version'])
                if (kind,ident) in seen: raise ValueError('Record ID appears more than once in this upload.')
                seen.add((kind,ident))
                row={'kind':kind,'id':ident,'expected_version':version,'document':doc,'data':delivery(val) if kind=='delivery' else charge(val),'sheet':name,'row_number':index,'source_payload':clean(val)}
                if kind=='delivery' and _blank(val['charge_id']) and not _blank(val['charge_version']):
                    raise ValueError('A charge version without its charge ID is invalid.')
                if kind=='delivery' and (not _blank(val['charge_id']) or any(not _blank(val[f]) for f in CHARGE_FIELDS)):
                    cid,cv=record_id(val['charge_id'],val['charge_version'])
                    if ('charge',cid) in seen: raise ValueError('Charge ID appears more than once in this upload.')
                    seen.add(('charge',cid)); row['charge']={'id':cid,'expected_version':cv,'data':charge(val)}
                result['rows'].append(row)
            except ValueError as exc:
                result['issues'].append(issue('INVALID_RECORD',str(exc),name,[f'A{index}']))
        result['periods'].append({'sheet':name,'period':'Stable-ID records','do_rows':count if kind=='delivery' else 0})

def _read_legacy(book,cached,result):
    repeated=Counter()
    for sheet in book:
        sheet.reset_dimensions(); cache=cached[sheet.title]; cache.reset_dimensions()
        rows=list(sheet.iter_rows(max_col=64)); cr=list(cache.iter_rows(max_col=64))
        get=lambda r,c: rows[r-1][c-1].value if r<=len(rows) else None
        if tuple(_normalize(get(6,c)) for c in range(1,20))!=tuple(_normalize(x) for x in HEADERS):
            result['issues'].append(issue('HEADER_MISMATCH','Expected the reviewed 19-column layout in row 6.',sheet.title,['A6:S6'])); continue
        period=_period(get(3,3)); invoice=text_value(get(4,3))
        if not period or not invoice:
            result['issues'].append(issue('INVALID_DOCUMENT','Invoice reference and recognized billing period are required.',sheet.title,['C3:C4'])); continue
        doc={'invoice_number':invoice,'period_start':period[1].isoformat(),'period_end':period[2].isoformat()}
        count=0; skipped=0; aux=0; blanks=0
        for index,cells in enumerate(rows,1):
            if index<7: continue
            raw=[c.value for c in cells[:19]]
            aux+=int(any(not _blank(c.value) for c in cells[19:]))
            if all(_blank(v) for v in raw): continue
            if isinstance(raw[0],str) and ' '.join(raw[0].upper().split()) in SUMMARY_LABELS:
                skipped+=1; continue
            try:
                standalone=all(_blank(v) for v in raw[:14]) and any(not _blank(raw[i]) for i in (14,15,16))
                vals=list(raw)
                for col,c in enumerate(cells[:19]):
                    if c.data_type!='f': continue
                    # Only the exact row sum used by this legacy layout is accepted.
                    formula=re.sub(r'\s+','',str(c.value)).upper()
                    if col!=16 or formula not in {f'=SUM(O{index}+P{index})',f'=SUM(O{index}:P{index})',f'=O{index}+P{index}'}:
                        raise ValueError('Unsupported formula; replace it with a reviewed value.')
                    a=number(raw[14],money=True); b=number(raw[15],money=True)
                    computed=Decimal(a or '0')+Decimal(b or '0')
                    cached_value=cr[index-1][col].value if index<=len(cr) else None
                    if cached_value is not None and Decimal(str(cached_value))!=computed: raise ValueError('Stored total disagrees with the approved row sum.')
                    vals[col]=str(computed)
                row={'kind':'charge' if standalone else 'delivery','id':str(uuid4()),'expected_version':None,'document':doc,'sheet':sheet.title,'row_number':index,'source_payload':dict(zip(HEADERS,clean(raw)))}
                cv=charge(dict(zip(CHARGE_FIELDS,vals[14:18])))
                if standalone:
                    row['data']=cv
                    result['issues'].append(issue('INVOICE_LEVEL_CHARGE','Standalone charge retained at invoice level; no delivery or trip was inferred.',sheet.title,[f'O{index}:R{index}'],'warning'))
                else:
                    mapped=dict(zip(DELIVERY_FIELDS, [vals[2],vals[3],vals[1],vals[4],vals[5],vals[6],vals[7],vals[8],vals[9],vals[10],vals[11],vals[12],vals[13],vals[17]]))
                    row['data']=delivery(mapped); count+=1; repeated[row['data']['do_number'].upper()]+=1
                    if isinstance(raw[2],(int,float)): result['issues'].append(issue('NUMERIC_ID','Numeric D/O preserved as text; previously lost leading zeros cannot be reconstructed.',sheet.title,[f'C{index}'],'warning'))
                    blanks+=int(any(vals[i] is None for i in (1,10,11,13,14,16)))
                    if any(not _blank(vals[i]) for i in (14,15,16)):
                        row['charge']={'id':str(uuid4()),'expected_version':None,'data':cv}
                result['rows'].append(row)
            except ValueError as exc:
                result['issues'].append(issue('INVALID_RECORD',str(exc),sheet.title,[f'A{index}:S{index}']))
        result['periods'].append({'sheet':sheet.title,'period':period[0],'period_start':doc['period_start'],'period_end':doc['period_end'],'do_rows':count})
        if blanks: result['issues'].append(issue('NULLS_PRESERVED',f'{blanks} delivery rows contain optional/grouped blanks. Blanks remain null; no values or charges were filled down.',sheet.title,severity='warning'))
        result['issues'].append(issue('REFERENCE_ROWS',f'{skipped} recognized summary rows and {aux} rows with auxiliary reference columns were excluded from operational records; their source remains in the workbook.',sheet.title,severity='info'))
    if any(n>1 for n in repeated.values()):
        result['issues'].append(issue('REPEATED_DO_PRESERVED','Repeated D/O references are preserved as separate source rows with independent stable IDs, not used as update keys.',severity='warning'))
