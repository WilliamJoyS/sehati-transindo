"""Customer-scoped staged imports, atomic version-checked confirmation and export."""
from __future__ import annotations
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
import re
import openpyxl
from openpyxl.styles import Font, PatternFill
from psycopg import sql
from psycopg.types.json import Jsonb
from .db import connect
from .workbook import WorkbookValidationError
from .import_engine import (read_workbook, clean, issue, VERSION, LEGACY_VERSION,
                            DELIVERY_FIELDS, CHARGE_FIELDS, DOCUMENT_FIELDS,
                            DELIVERY_HEADERS, CHARGE_HEADERS)

class OperationsError(ValueError):
    def __init__(self,status_code,message):
        self.status_code=status_code; self.message=message
        super().__init__(message)

def _uuid(value):
    try: return str(UUID(str(value)))
    except (ValueError,TypeError): raise OperationsError(422,'Invalid identifier.')

def _audit(conn,actor,action,object_id,details=None):
    conn.execute('INSERT INTO audit_events(actor_id,action,object_id,details) VALUES(%s,%s,%s,%s)',(actor,action,str(object_id),Jsonb(clean(details or {}))))

def _customer(conn,customer_id,lock=False):
    row=conn.execute('SELECT * FROM customers WHERE id=%s'+(' FOR UPDATE' if lock else ''),(_uuid(customer_id),)).fetchone()
    if not row or not row['active']: raise OperationsError(404,'Active customer not found.')
    return row

def create_customer(data,actor_id=None):
    code=str(data.get('code','')).strip().upper(); name=str(data.get('name','')).strip()
    if not re.fullmatch(r'[A-Z0-9][A-Z0-9_-]{1,39}',code) or not 2<=len(name)<=200:
        raise OperationsError(422,'Customer requires a 2–40 character code and a 2–200 character name.')
    with connect() as conn:
        row=conn.execute('INSERT INTO customers(id,code,name) VALUES(%s,%s,%s) ON CONFLICT(code) DO NOTHING RETURNING *',(uuid4(),code,name)).fetchone()
        if row is None: raise OperationsError(409,'Customer code already exists.')
        _audit(conn,actor_id,'customer_created',row['id'],{'code':code})
        return clean(row)

def list_customers():
    with connect() as conn: return clean(conn.execute('SELECT * FROM customers ORDER BY name').fetchall())

def _fields(row,fields): return {f:clean(row.get(f)) for f in fields}

def _change_list(before,after):
    return [{'field':f,'old':(before or {}).get(f),'new':v} for f,v in after.items() if (before or {}).get(f)!=v]

def _classify(conn,staged,customer_id):
    """Read current state; caller locks the customer during confirmation."""
    errors=[]; known_docs={}; legacy=staged['mapping_version']==LEGACY_VERSION
    for name,doc in staged['documents'].items():
        existing=conn.execute('SELECT * FROM billing_documents WHERE customer_id=%s AND invoice_number=%s',(customer_id,name)).fetchone()
        info={'id':str(existing['id']) if existing else doc.setdefault('id',str(uuid4())),'exists':bool(existing),'unchanged':False,'blocked':False}
        if existing and any(clean(existing[k])!=doc[k] for k in ('period_start','period_end')):
            errors.append(issue('DOCUMENT_PERIOD_CONFLICT','Existing invoice has a different billing period. Review its reference rather than rewriting historical document metadata.'))
            info['blocked']=True
        if existing and legacy:
            match=conn.execute('SELECT 1 FROM legacy_document_snapshots WHERE customer_id=%s AND invoice_number=%s AND content_digest=%s',(customer_id,name,doc['content_digest'])).fetchone()
            if match: info['unchanged']=True
            else:
                info['blocked']=True
                errors.append(issue('LEGACY_UPDATE_REQUIRES_IDS','A legacy invoice already exists with different contents. Download its stable-ID export and apply corrections there; D/O and row positions are not update keys.'))
        known_docs[name]=info; doc['id']=info['id']
    for row in staged['rows']:
        doc=known_docs[row['document']['invoice_number']]
        row.update(action='new',changes=[],before=None)
        if doc['blocked']: row['action']='blocked'; continue
        if doc['unchanged']:
            row['action']='unchanged'
            row.pop('document_id',None)
            continue
        row['document_id']=doc['id']
        table='deliveries' if row['kind']=='delivery' else 'invoice_lines'
        fields=DELIVERY_FIELDS if row['kind']=='delivery' else CHARGE_FIELDS
        old=conn.execute(sql.SQL('SELECT * FROM {} WHERE id=%s').format(sql.Identifier(table)),(row['id'],)).fetchone()
        if old:
            if str(old['customer_id'])!=str(customer_id) or row['expected_version']!=old['version']:
                row['action']='blocked'; errors.append(issue('STALE_OR_FOREIGN_RECORD','Record is stale or does not belong to this customer. Export current data and review the changes.',row['sheet'],[str(row['row_number'])])); continue
            if row['kind']=='charge' and old['delivery_id'] is not None:
                row['action']='blocked'; errors.append(issue('CHARGE_LINK_MISMATCH','Delivery-linked charges must remain on their Deliveries row.',row['sheet'],[str(row['row_number'])])); continue
            before=_fields(old,fields); row['before']=clean(old); row['changes']=_change_list(before,row['data'])
            doc_field='source_document_id' if row['kind']=='delivery' else 'billing_document_id'
            if str(old[doc_field])!=doc['id']: row['changes'].append({'field':'invoice_number','old':'previous document','new':row['document']['invoice_number']})
            row['action']='updated' if row['changes'] else 'unchanged'
        elif row['expected_version'] is not None:
            row['action']='blocked'; errors.append(issue('UNKNOWN_RECORD','Record ID no longer exists; existing IDs cannot create new records.',row['sheet'],[str(row['row_number'])])); continue
        else: row['changes']=_change_list(None,row['data'])
        if row['kind']!='delivery': continue
        linked=conn.execute('SELECT * FROM invoice_lines WHERE delivery_id=%s',(row['id'],)).fetchone()
        charge=row.get('charge')
        if not charge:
            if linked:
                row['action']='blocked'; errors.append(issue('MISSING_CHARGE_ID','Existing charge identity is missing. Keep its exported ID/version, even when clearing amounts.',row['sheet'],[str(row['row_number'])]))
            continue
        oldcharge=conn.execute('SELECT * FROM invoice_lines WHERE id=%s',(charge['id'],)).fetchone()
        if oldcharge:
            if (str(oldcharge['customer_id'])!=str(customer_id) or str(oldcharge['delivery_id'])!=row['id'] or charge['expected_version']!=oldcharge['version']):
                row['action']='blocked'; errors.append(issue('CHARGE_CONFLICT','Charge identity, ownership, or version changed. Export current records before editing.',row['sheet'],[str(row['row_number'])])); continue
            charge['before']=clean(oldcharge); charge['changes']=_change_list(_fields(oldcharge,CHARGE_FIELDS),charge['data'])
            charge['action']='updated' if charge['changes'] or str(oldcharge['billing_document_id'])!=doc['id'] else 'unchanged'
        elif charge['expected_version'] is not None or linked:
            row['action']='blocked'; errors.append(issue('CHARGE_CONFLICT','Charge identity is missing or conflicts with an existing delivery charge.',row['sheet'],[str(row['row_number'])])); continue
        else: charge.update(action='new',before=None,changes=_change_list(None,charge['data']))
        if charge['action']!='unchanged' and row['action']=='unchanged': row['action']='updated'
        row['changes'].extend({**x,'field':'charge.'+x['field']} for x in charge['changes'])
    return errors

def _report(staged,batch_id,customer_id,filename,file_digest,revision=1,status=None):
    rows=staged['rows']; blocking=any(x['severity']=='blocking' for x in staged['issues'])
    actions={key:sum(r.get('action')==key for r in rows) for key in ('new','updated','unchanged','blocked')}
    report={'batch_id':str(batch_id),'customer_id':str(customer_id),'source':filename,'revision':revision,
        'status':status or ('blocked' if blocking else 'preview'),'mapping_version':staged['mapping_version'],
        'source_metadata':{'sha256':file_digest,'mapping_version':staged['mapping_version']},
        'summary':{'sheets':len(staged['periods']),'do_rows':sum(r['kind']=='delivery' for r in rows),
                   'distinct_do':len({r['data']['do_number'].upper() for r in rows if r['kind']=='delivery'}),
                   'charges':sum(r['kind']=='charge' or bool(r.get('charge')) for r in rows),
                   'issue_count':len(staged['issues']),**actions},
        'periods':staged['periods'],'issues':staged['issues'],
        'can_confirm':not blocking and status!='committed','commit_enabled':not blocking and status!='committed',
        'rows':[{'source_sheet':r['sheet'],'source_row':r['row_number'],'record_id':r['id'],
                 'do_number':r['data'].get('do_number'),'kind':r['kind'],'action':r.get('action','blocked'),'changes':r.get('changes',[])} for r in rows[:50]],
        'total_stage_rows':len(rows),'truncated':len(rows)>50}
    return report

def _committed_replay(conn,customer_id,file_digest,staged):
    return conn.execute("""SELECT id,report,staged FROM import_batches
        WHERE customer_id=%s AND mapping_version=%s AND status='committed'
        AND (file_digest=%s OR staged->>'content_digest'=%s)
        ORDER BY committed_at DESC LIMIT 1""",
        (customer_id,staged['mapping_version'],file_digest,staged.get('content_digest'))).fetchone()


def stage_import(path,filename,customer_id,actor_id=None):
    customer_id=_uuid(customer_id); filename=Path(filename.replace('\\','/')).name[:255]
    file_digest=sha256(Path(path).read_bytes()).hexdigest()
    with connect() as conn:
        _customer(conn,customer_id)
    try:
        staged=read_workbook(path,customer_id)
    except WorkbookValidationError as exc:
        batch_id=str(uuid4())
        staged={'mapping_version':'unsupported','rows':[],'documents':{},'periods':[],
                'issues':[issue('UPLOAD_REJECTED',str(exc))]}
        report=_report(staged,batch_id,customer_id,filename,file_digest,status='failed')
        report.update(can_confirm=False,commit_enabled=False)
        with connect() as conn:
            conn.execute('INSERT INTO import_batches(id,customer_id,staff_id,filename,file_digest,mapping_version,status,report,staged) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)',
                         (batch_id,customer_id,actor_id,filename,file_digest,'unsupported','failed',Jsonb(report),Jsonb(staged)))
            _audit(conn,actor_id,'import_rejected',batch_id,{'reason':str(exc)})
        raise
    with connect() as conn:
        _customer(conn,customer_id)
        prior=None if any(x['severity']=='blocking' for x in staged['issues']) else _committed_replay(conn,customer_id,file_digest,staged)
        if prior:
            report=prior['report']; report['reused']=True
            _audit(conn,actor_id,'import_replayed',report['batch_id'],{'customer_id':customer_id})
            return report
        staged['issues'].extend(_classify(conn,staged,customer_id))
        batch_id=str(uuid4()); report=_report(staged,batch_id,customer_id,filename,file_digest)
        conn.execute('INSERT INTO import_batches(id,customer_id,staff_id,filename,file_digest,mapping_version,status,report,staged) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)',
                     (batch_id,customer_id,actor_id,filename,file_digest,staged['mapping_version'],report['status'],Jsonb(report),Jsonb(staged)))
        _audit(conn,actor_id,'import_staged',batch_id,{'customer_id':customer_id,'summary':report['summary']})
        return report

def get_batch(batch_id):
    with connect() as conn:
        row=conn.execute('SELECT report FROM import_batches WHERE id=%s',(_uuid(batch_id),)).fetchone()
        if not row: raise OperationsError(404,'Import batch not found.')
        return row['report']


def get_batch_rows(batch_id,limit=50,offset=0):
    with connect() as conn:
        row=conn.execute('SELECT staged,report FROM import_batches WHERE id=%s',(_uuid(batch_id),)).fetchone()
        if not row: raise OperationsError(404,'Import batch not found.')
    rows=row['staged']['rows']
    reused=row['report'].get('reused',False)
    return {'items':[{'source_sheet':r['sheet'],'source_row':r['row_number'],
                     'record_id':r['id'],'do_number':r['data'].get('do_number'),
                     'kind':r['kind'],'action':'unchanged' if reused else r.get('action','blocked'),
                     'changes':[] if reused else r.get('changes',[])} for r in rows[offset:offset+limit]],
            'total':len(rows),'limit':limit,'offset':offset}

def list_batches(customer_id=None):
    with connect() as conn:
        where=' WHERE customer_id=%s' if customer_id else ''; params=(_uuid(customer_id),) if customer_id else ()
        rows=conn.execute('SELECT id,customer_id,filename,status,revision,created_at,committed_at,report->\'summary\' AS summary FROM import_batches'+where+' ORDER BY created_at DESC LIMIT 100',params).fetchall()
        return clean(rows)

def _write(conn,table,record_id,customer_id,document_id,data,fields,action,delivery_id=None):
    docfield='source_document_id' if table=='deliveries' else 'billing_document_id'
    values={'customer_id':customer_id,docfield:document_id,**data}
    if table=='invoice_lines': values['delivery_id']=delivery_id
    if action=='new':
        values={'id':record_id,**values}
        conn.execute(sql.SQL('INSERT INTO {} ({}) VALUES ({})').format(sql.Identifier(table),sql.SQL(',').join(map(sql.Identifier,values)),sql.SQL(',').join(sql.Placeholder() for _ in values)),tuple(values.values()))
    elif action=='updated':
        conn.execute(sql.SQL('UPDATE {} SET {},version=version+1,updated_at=now() WHERE id=%s').format(sql.Identifier(table),sql.SQL(',').join(sql.SQL('{}=%s').format(sql.Identifier(k)) for k in values)),tuple(values.values())+(record_id,))

def confirm_import(batch_id,actor_id=None,expected_revision=None):
    batch_id=_uuid(batch_id)
    try:
        with connect() as conn:
            # One customer lock serializes all its imports and partner data revisions.
            info=conn.execute('SELECT customer_id FROM import_batches WHERE id=%s',(batch_id,)).fetchone()
            if not info: raise OperationsError(404,'Import batch not found.')
            customer=_customer(conn,info['customer_id'],lock=True)
            batch=conn.execute('SELECT * FROM import_batches WHERE id=%s FOR UPDATE',(batch_id,)).fetchone()
            if batch['status']=='committed': return {**batch['report'],'reused':True}
            if batch['status']!='preview': raise OperationsError(409,'Batch is blocked or failed; correct the file and upload a fresh preview.')
            if expected_revision is not None and expected_revision!=batch['revision']: raise OperationsError(409,'Preview revision changed; refresh the preview.')
            staged=batch['staged']
            # Two concurrent previews must not commit the same new rows twice.
            # The customer lock above serializes this replay check with writes.
            prior=_committed_replay(conn,batch['customer_id'],batch['file_digest'],staged)
            if prior:
                report=deepcopy(prior['report'])
                report.update(batch_id=batch_id,reused=True,reused_batch_id=str(prior['id']))
                report['summary'].update(new=0,updated=0,unchanged=len(staged['rows']))
                report['rows']=[{**r,'action':'unchanged','changes':[]} for r in report['rows']]
                conn.execute("UPDATE import_batches SET status='committed',report=%s,staged=%s,committed_at=now() WHERE id=%s",(Jsonb(report),Jsonb(prior['staged']),batch_id))
                _audit(conn,actor_id,'import_replayed',batch_id,{'reused_batch_id':str(prior['id'])})
                return report
            errors=_classify(conn,staged,str(batch['customer_id']))
            if errors: raise OperationsError(409,'Database changes conflict with this preview. Nothing was imported. Download current records and create a fresh preview.')
            changed=any(r['action'] in ('new','updated') for r in staged['rows'])
            for doc in staged['documents'].values():
                conn.execute('INSERT INTO billing_documents(id,customer_id,invoice_number,period_start,period_end) VALUES(%s,%s,%s,%s,%s) ON CONFLICT(customer_id,invoice_number) DO NOTHING',(doc['id'],batch['customer_id'],doc['invoice_number'],doc['period_start'],doc['period_end']))
            for row in staged['rows']:
                action=row['action']; fields=DELIVERY_FIELDS if row['kind']=='delivery' else CHARGE_FIELDS
                if action in ('new','updated'):
                    _write(conn,'deliveries' if row['kind']=='delivery' else 'invoice_lines',row['id'],batch['customer_id'],row['document_id'],row['data'],fields,action)
                    if row.get('charge'):
                        ch=row['charge']; _write(conn,'invoice_lines',ch['id'],batch['customer_id'],row['document_id'],ch['data'],CHARGE_FIELDS,ch['action'],row['id'])
                conn.execute('INSERT INTO import_rows(batch_id,sheet,row_number,kind,record_id,action,source_payload,before_payload,after_payload) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)',
                    (batch_id,row['sheet'],row['row_number'],row['kind'],row['id'] if row.get('document_id') else None,action,Jsonb(row['source_payload']),Jsonb(row.get('before')),Jsonb({'data':row['data'],'charge':row.get('charge')})))
            if staged['mapping_version']==LEGACY_VERSION:
                for doc in staged['documents'].values():
                    conn.execute('INSERT INTO legacy_document_snapshots(customer_id,invoice_number,content_digest,batch_id) VALUES(%s,%s,%s,%s) ON CONFLICT DO NOTHING',(batch['customer_id'],doc['invoice_number'],doc['content_digest'],batch_id))
            if changed: conn.execute('UPDATE customers SET data_version=data_version+1 WHERE id=%s',(batch['customer_id'],))
            report=_report(staged,batch_id,batch['customer_id'],batch['filename'],batch['file_digest'],batch['revision'],status='committed')
            report['customer_data_version']=customer['data_version']+int(changed)
            conn.execute("UPDATE import_batches SET status='committed',report=%s,staged=%s,committed_at=now() WHERE id=%s",(Jsonb(report),Jsonb(staged),batch_id))
            _audit(conn,actor_id,'import_committed',batch_id,{'customer_id':str(batch['customer_id']),'summary':report['summary'],'customer_data_version':report['customer_data_version']})
            return report
    except OperationsError as exc:
        # Separate transaction preserves the failure event after business rollback.
        if exc.status_code==409:
            with connect() as conn:
                row=conn.execute('SELECT report FROM import_batches WHERE id=%s FOR UPDATE',(batch_id,)).fetchone()
                if row:
                    report=row['report']
                    if report['status']=='preview':
                        report.update(status='failed',can_confirm=False,commit_enabled=False)
                        report['issues'].append(issue('CONFIRM_CONFLICT',exc.message)); report['summary']['issue_count']=len(report['issues'])
                        conn.execute("UPDATE import_batches SET status='failed',report=%s WHERE id=%s",(Jsonb(report),batch_id))
                    _audit(conn,actor_id,'import_confirm_failed',batch_id,{'reason':exc.message})
        raise
    except Exception as exc:
        with connect() as conn:
            row=conn.execute('SELECT report FROM import_batches WHERE id=%s FOR UPDATE',(batch_id,)).fetchone()
            if row and row['report']['status']=='preview':
                report=row['report']; report.update(status='failed',can_confirm=False,commit_enabled=False)
                report['issues'].append(issue('COMMIT_FAILED','Transaction failed; no business records were changed. Create a fresh preview.'))
                report['summary']['issue_count']=len(report['issues'])
                conn.execute("UPDATE import_batches SET status='failed',report=%s WHERE id=%s",(Jsonb(report),batch_id))
            _audit(conn,actor_id,'import_confirm_failed',batch_id,{'reason':'Database transaction failed and rolled back.'})
        raise OperationsError(409,'Import transaction failed and rolled back. No business changes were committed.') from exc

def list_deliveries(customer_id=None,limit=50,offset=0):
    limit=max(1,min(int(limit),200)); offset=max(0,int(offset))
    where=' WHERE d.customer_id=%s' if customer_id else ''; params=(_uuid(customer_id),) if customer_id else ()
    with connect() as conn:
        total=conn.execute('SELECT count(*) AS n FROM deliveries d'+where,params).fetchone()['n']
        rows=conn.execute('SELECT d.*,b.invoice_number,b.period_start,b.period_end,c.name AS customer_name FROM deliveries d JOIN billing_documents b ON b.id=d.source_document_id JOIN customers c ON c.id=d.customer_id'+where+' ORDER BY d.load_date DESC,d.id LIMIT %s OFFSET %s',params+(limit,offset)).fetchall()
        return {'items':clean(rows),'total':total,'limit':limit,'offset':offset}

def export_workbook(customer_id):
    customer_id=_uuid(customer_id)
    with connect() as conn:
        conn.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY')
        customer=_customer(conn,customer_id)
        deliveries=conn.execute('SELECT d.*,b.invoice_number,b.period_start,b.period_end FROM deliveries d JOIN billing_documents b ON b.id=d.source_document_id WHERE d.customer_id=%s ORDER BY b.period_start,d.created_at,d.id',(customer_id,)).fetchall()
        charges=conn.execute('SELECT l.*,b.invoice_number,b.period_start,b.period_end FROM invoice_lines l JOIN billing_documents b ON b.id=l.billing_document_id WHERE l.customer_id=%s ORDER BY l.created_at,l.id',(customer_id,)).fetchall()
    if len(deliveries)+sum(c['delivery_id'] is None for c in charges)>10000:
        raise OperationsError(422,'Export exceeds the current 10,000 row workbook limit; use a scoped export before increasing limits.')
    book=openpyxl.Workbook(); meta=book.active; meta.title='Metadata'
    meta.append(['template_version',VERSION]); meta.append(['customer_id',customer_id]); meta.append(['customer_name',customer['name']]); meta.append(['data_version',customer['data_version']])
    meta.append(['instructions','Keep record IDs and versions. New records: leave both blank. Missing rows never delete. Blank editable cells clear that field after preview. Values only; no formulas.'])
    linked={str(c['delivery_id']):c for c in charges if c['delivery_id']}
    ds=book.create_sheet('Deliveries'); ds.append(DELIVERY_HEADERS)
    for row in deliveries:
        charge=linked.get(str(row['id']),{})
        values={'record_id':str(row['id']),'record_version':row['version'],**_fields(row,DOCUMENT_FIELDS+DELIVERY_FIELDS),
                'charge_id':str(charge['id']) if charge else None,'charge_version':charge.get('version'),**_fields(charge,CHARGE_FIELDS)}
        ds.append([values.get(k) for k in DELIVERY_HEADERS])
    cs=book.create_sheet('Charges'); cs.append(CHARGE_HEADERS)
    for row in charges:
        if row['delivery_id']: continue
        values={'charge_id':str(row['id']),'charge_version':row['version'],**_fields(row,DOCUMENT_FIELDS+CHARGE_FIELDS)}
        cs.append([values.get(k) for k in CHARGE_HEADERS])
    for sheet in book:
        sheet.freeze_panes='A2'
        if sheet.title!='Metadata': sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='234D3C')
        for row in sheet:
            for cell in row:
                # Force text cells to remain text, including D/O or notes beginning '='.
                if isinstance(cell.value,str): cell.data_type='s'
        for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(35,max(15,len(str(col[0].value or ''))+2))
    output=BytesIO(); book.save(output); return output.getvalue()
